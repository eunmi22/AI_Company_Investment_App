"""3차 고도화: 뉴스 관련성 필터·AI 분류·집계·Excel 보고서 모의 테스트.

실제 API 키나 뉴스 API 호출 없이 실행합니다.
실행: python -m unittest test_news_analysis -v
"""
import json
import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pandas as pd
from openpyxl import load_workbook

from news_analysis import (
    _canonical_url, _company_relevance, analyze_news_articles,
    prepare_news, sentiment_summary, verify_news_analysis,
)
from core import excel_bytes, make_ai_context, report_markdown


class TestNewsAnalysis(unittest.TestCase):
    def setUp(self):
        self.news = pd.DataFrame([
            {'발행일':'2026-10-20T07:00:00Z', '제목':'삼성전자 차세대 메모리 제품 공개',
             '언론사':'예시신문', '요약':'삼성전자가 차세대 제품을 공개했다.',
             '링크':'https://example.com/a?utm_source=x', '출처':'테스트'},
            {'발행일':'2026-10-20T06:00:00Z', '제목':'삼성전자 실적 전망 하향 발표',
             '언론사':'경제신문', '요약':'시장 전망 관련 소식',
             '링크':'https://example.com/b', '출처':'테스트'},
            {'발행일':'2026-10-20T05:00:00Z', '제목':'서울 환경단체 정화 사업 실시',
             '언론사':'지역신문', '요약':'삼성전자 언급 없음. 사실 이 요약에는 삼성전자가 등장합니다.',
             '링크':'https://example.com/c', '출처':'테스트'},
            {'발행일':'2026-10-20T04:00:00Z', '제목':'래블업 GPU 인프라 사업 확대',
             '언론사':'기술신문', '요약':'상장 계획 논의',
             '링크':'https://example.com/d', '출처':'테스트'},
        ])

    def mock_client(self, news, invalid=False, related='직접'):
        """입력 기사 ID에 맞춰 결정적인 가짜 JSON을 반환합니다."""
        client = MagicMock()
        def complete(**kwargs):
            user = json.loads(kwargs['messages'][1]['content'])
            items = []
            for record in user['articles']:
                sid = record['id']
                tag = '부정' if '하향' in record['제목'] else '긍정' if '메모리' in record['제목'] else '중립'
                items.append({'id': ('잘못된ID' if invalid else sid), '감성': tag,
                              'AI관련성': related, '영향분야': '실적·재무' if tag == '부정' else '제품·기술',
                              '위험요인': '실적 둔화 우려' if tag == '부정' else '확인 필요',
                              '판단근거': '제목에 제품 공개 또는 실적 하향 정보가 보임'})
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({'items':items}, ensure_ascii=False)))])
        client.chat.completions.create.side_effect = complete
        return client

    def test_relevance(self):
        self.assertEqual(_company_relevance('삼성전자', '삼성전자 신제품 출시', ''), '직접')
        self.assertEqual(_company_relevance('삼성전자', '삼성 갤럭시 신제품 출시', ''), '추정')
        self.assertEqual(_company_relevance('삼성전자', '삼성디스플레이 OLED 공장', ''), '무관')
        self.assertEqual(_company_relevance('삼성전자', '래블업 IPO', ''), '무관')

    def test_unrelated_removed_and_keeps_traced(self):
        selected, excluded = prepare_news(self.news, '삼성전자')
        self.assertEqual(len(selected), 3)
        self.assertEqual(len(excluded), 1)
        self.assertIn('래블업', excluded.iloc[0]['제목'])
        self.assertEqual(list(selected['기사ID']), ['N001', 'N002', 'N003'])

    def test_duplicate_link_tracking(self):
        f = self.news.iloc[[0,0]].copy().reset_index(drop=True)
        f.loc[1, '링크'] = 'https://example.com/a?utm_source=y'
        selected, excluded = prepare_news(f, '삼성전자')
        self.assertEqual(len(selected), 1)
        self.assertEqual(len(excluded), 1)
        self.assertEqual(_canonical_url('https://abc.com/news?utm_source=y&id=5'), 'https://abc.com/news?id=5')

    def test_near_duplicate_different_numeric_is_not_collapsed(self):
        f = self.news.iloc[[0]].copy()
        f = pd.concat([f, f.copy().assign(링크='https://example.com/unique')], ignore_index=True)
        selected, excluded = prepare_news(f, '삼성전자')
        self.assertEqual(len(selected), 1)
        self.assertIn('유사 제목 중복', excluded.iloc[0]['제외사유'])
        f.loc[1, '제목'] = '삼성전자 2조원 수주'
        f.loc[0, '제목'] = '삼성전자 3조원 수주'
        selected, _ = prepare_news(f, '삼성전자')
        self.assertEqual(len(selected), 2)

    def test_invalid_link_and_html(self):
        f = self.news.iloc[[0]].copy()
        f.loc[f.index[0], '제목'] = '&quot;<b>삼성전자</b> 신제품&quot;'
        selected, _ = prepare_news(f, '삼성전자')
        self.assertNotIn('<b>', selected.iloc[0]['제목'])
        f.loc[f.index[0], '링크'] = 'javascript:alert(1)'
        selected, excluded = prepare_news(f, '삼성전자')
        self.assertEqual(len(selected), 0)
        self.assertEqual(len(excluded), 1)

    def test_only_company_specific_inferences(self):
        f = pd.DataFrame([{'발행일':'2026-10-19', '제목':'삼성 폴더블 갤럭시 신규 모델',
                           '언론사':'테스트', '요약':'삼성 브랜드 제품',
                           '링크':'https://example.com/galaxy', '출처':'테스트'}])
        selected, excluded = prepare_news(f, '삼성전자')
        self.assertEqual(selected.iloc[0]['사전관련성'], '추정')
        self.assertTrue(excluded.empty)

    def test_max_article_limit(self):
        f = pd.concat([self.news.iloc[[0]].assign(제목=f'삼성전자 발표 {i}년', 링크=f'https://example.com/item/{i}')
                       for i in range(6)], ignore_index=True)
        selected, excluded = prepare_news(f, '삼성전자', limit=3)
        self.assertEqual(len(selected), 3)
        self.assertEqual(len(excluded), 3)
        self.assertTrue(excluded['제외사유'].str.contains('표시 기사 수 제한').all())

    def test_empty_no_api(self):
        empty, excluded = prepare_news(pd.DataFrame(), '삼성전자')
        analyzed = analyze_news_articles(empty, '삼성전자', openai_key='')
        self.assertTrue(analyzed.empty)
        self.assertTrue(excluded.empty)

    def test_api_key_required(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        with self.assertRaisesRegex(ValueError, 'OPENAI_API_KEY'):
            analyze_news_articles(selected, '삼성전자', openai_key='')

    def test_sentiment_and_risk(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        client = self.mock_client(selected)
        analyzed = analyze_news_articles(selected, '삼성전자', openai_key='example', client=client)
        self.assertEqual(list(analyzed['감성']), ['긍정', '부정', '중립'])
        self.assertIn('실적 둔화', analyzed.iloc[1]['위험요인'])
        self.assertEqual(analyzed.iloc[1]['링크'], selected.iloc[1]['링크'])
        self.assertEqual(analyzed.iloc[1]['분석상태'], 'AI 추정 · 원문 검증 필요')
        self.assertEqual(client.chat.completions.create.call_count, 1)
        data = client.chat.completions.create.call_args.kwargs
        self.assertEqual(data['response_format']['type'], 'json_object')
        self.assertIn('신뢰되지 않는', data['messages'][0]['content'])

    def test_wrong_id_is_rejected(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        with self.assertRaisesRegex(ValueError, '기사ID'):
            analyze_news_articles(selected, '삼성전자', openai_key='x', client=self.mock_client(selected, invalid=True))

    def test_bad_sentiment_is_rejected(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        client = self.mock_client(selected)
        orig = client.chat.completions.create.side_effect
        def bad(**kwargs):
            response = orig(**kwargs)
            content = json.loads(response.choices[0].message.content)
            content['items'][0]['감성'] = '급등'
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(content)))])
        client.chat.completions.create.side_effect = bad
        with self.assertRaisesRegex(ValueError, '감성/관련성'):
            analyze_news_articles(selected, '삼성전자', openai_key='x', client=client)

    def test_irrelevance_is_not_counted(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        analyzed = analyze_news_articles(selected, '삼성전자', openai_key='x',
                                         client=self.mock_client(selected, related='무관'))
        self.assertTrue((analyzed['감성'] == '판단보류').all())
        counts = sentiment_summary(analyzed)
        self.assertEqual(int(counts['기사수'].sum()), 0)

    def test_summary_percent_and_verification(self):
        selected, _ = prepare_news(self.news, '삼성전자')
        analyzed = analyze_news_articles(selected, '삼성전자', openai_key='x', client=self.mock_client(selected))
        counts = sentiment_summary(analyzed)
        self.assertEqual(int(counts['기사수'].sum()), len(selected))
        self.assertAlmostEqual(float(counts['비중(%)'].sum()), 100.0, places=1)
        checks = verify_news_analysis(selected, analyzed)
        self.assertTrue((checks.iloc[:3]['상태'] == '확인').all())
        self.assertEqual(checks.iloc[-1]['상태'], '확인 필요')
        tampered = analyzed.copy()
        tampered.loc[0, '링크'] = 'https://fakedomain.example/false'
        chk = verify_news_analysis(selected, tampered)
        self.assertEqual(chk.loc[chk['검증항목']=='기사 원문 링크 보존','상태'].iloc[0], '불일치')

    def test_ai_batching_cost_bounded(self):
        df = pd.concat([self.news.iloc[[0]].assign(제목=f'삼성전자 발표 기사 {i}',
                                    링크=f'https://example.com/story/{i}') for i in range(12)], ignore_index=True)
        prepared, _ = prepare_news(df, '삼성전자')
        client = self.mock_client(prepared)
        analyzed = analyze_news_articles(prepared, '삼성전자', openai_key='x', client=client, batch_size=5)
        self.assertEqual(len(analyzed), 12)
        self.assertEqual(client.chat.completions.create.call_count, 3)
        with self.assertRaisesRegex(ValueError, '제한'):
            analyze_news_articles(prepared, '삼성전자', openai_key='x', client=client, max_articles=10)

    def test_export_report_context_has_analysis(self):
        selected, excluded = prepare_news(self.news, '삼성전자')
        analyzed = analyze_news_articles(selected, '삼성전자', openai_key='x', client=self.mock_client(selected))
        result = {
            '기업명': '삼성전자', '수집시각':'2026-10-20 09:00:00',
            '재무': pd.DataFrame(), '주가':pd.DataFrame(), '뉴스':selected,
            '뉴스분석': analyzed, '뉴스제외':excluded,
            '뉴스처리현황':pd.DataFrame([{'구분':'기사후보', '기사수':len(selected)}]),
            '품질':pd.DataFrame([{'자료':'뉴스'}]), '통계':pd.DataFrame(),
            '사실검증':pd.DataFrame([{'검증항목':'기사진위','상태':'확인 필요','근거/설명':'원문 미확인'}]),
            '출처':pd.DataFrame([{'자료':'뉴스','출처':'테스트'}]), '경고':[],
        }
        ctx = make_ai_context(result)
        self.assertEqual(ctx['뉴스검색결과(검색제목·요약·AI추정)'][0]['감성(기업영향추정)'], '긍정')
        wb = load_workbook(BytesIO(excel_bytes(result)))
        try:
            self.assertTrue({'뉴스감성분석','뉴스감성집계','뉴스감성검증','뉴스제외기록'}.issubset(wb.sheetnames))
            self.assertEqual(wb['뉴스감성분석'].max_row, len(selected) + 1)
        finally:
            wb.close()
        fake_report = {key: '원문 확인 필요' for key in
                       ('profitability', 'stability', 'trend', 'news_risks', 'conclusion')}
        md = report_markdown(result, fake_report)
        self.assertIn('뉴스 건별 AI 영향 분석', md)
        self.assertIn('실적 둔화 우려', md)


if __name__ == '__main__':
    unittest.main()
