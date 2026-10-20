"""외부 API/비용 없는 단위 검증. 실행: python -m unittest -v test_core.py"""
import unittest
from unittest.mock import patch, MagicMock
from types import ModuleType
import sys
import json
import pandas as pd
from core import (normalize_name, number_or_na, divide_percent,
                  items_from_body, choose_financial_record, parse_financial_row,
                  clean_and_quality, summarize_data, verify_facts, excel_bytes,
                  make_ai_context, report_markdown, ai_review)


class CoreTest(unittest.TestCase):
    def setUp(self):
        self.financial = pd.DataFrame([{
            '기업명': '테스트기업', '연도': 2024, '재무제표구분': '연결', '기준일자': '20241231',
            '매출액': 1000.0, '영업이익': 100.0, '당기순이익': 90.0,
            '총자산': 2000.0, '총부채': 800.0, '자기자본': 1200.0,
            '영업이익률(%)': 10.0, '순이익률(%)': 9.0, '부채비율(%)': 66.67, '매출증감률(%)': float('nan')
        }])
        self.prices = pd.DataFrame([{'날짜': pd.Timestamp('2024-12-01'), '시가': 10,
                                     '고가': 12, '저가': 9, '종가': 11, '거래량': 100, '출처': '테스트'}])
        self.news = pd.DataFrame([{'발행일': '2024-12-01', '제목': '테스트 기사', '언론사': '테스트뉴스',
                                   '요약': '', '링크': 'https://example.com/1', '출처': '테스트 RSS'}])
        self.result = {'기업명': '테스트기업', '종목코드': '000000', '수집시각': '2026-01-01 00:00:00',
                       '재무': self.financial, '주가': self.prices, '뉴스': self.news,
                       '품질': pd.DataFrame([{'자료': '재무', '원본행수': 1, '중복행수': 0, '결측셀수': 0, '정리후행수': 1}]),
                       '통계': pd.DataFrame([{'구분':'재무','지표':'매출액 평균','값':1000,'단위':'원','건수':1}]),
                       '사실검증': verify_facts('테스트기업', self.financial, self.prices, self.news, '000000'),
                       '출처': pd.DataFrame([{'자료':'재무', '출처':'테스트', '실제데이터수':1, '설명':'테스트'}]),
                       '경고': []}

    def test_number_math(self):
        self.assertEqual(divide_percent(100, 1000), 10.0)
        self.assertTrue(pd.isna(divide_percent(2, 0)))
        self.assertEqual(number_or_na('1,234'), 1234.0)
        self.assertTrue(pd.isna(number_or_na('-')))

    def test_company_normalize(self):
        self.assertEqual(normalize_name('삼성전자(주)'), normalize_name('삼성전자'))
        self.assertEqual(normalize_name('SK 하이닉스 주식회사'), normalize_name('SK하이닉스'))

    def test_api_items(self):
        self.assertEqual(items_from_body({'items': {'item': {'crno': '123'}}}), [{'crno': '123'}])
        self.assertEqual(items_from_body({'items': {'item': []}}), [])

    def test_financial_preference(self):
        rows = [{'fnclDcdNm': '별도재무제표', 'basDt': '20250101'},
                {'fnclDcdNm': '연결재무제표', 'basDt': '20241231'}]
        self.assertEqual(choose_financial_record(rows)['fnclDcdNm'], '연결재무제표')
        row = {'bizYear': '2024', 'fnclDcdNm': '연결재무제표', 'enpSaleAmt': '1,000',
               'enpBzopPft': '100', 'enpTdbtAmt':'300', 'enpTcptAmt':'600'}
        parsed = parse_financial_row(row, '테스트기업', 2024)
        self.assertEqual(parsed['영업이익률(%)'], 10.0)
        self.assertTrue(pd.isna(parsed['순이익률(%)']))
        self.assertEqual(parsed['부채비율(%)'], 50.0)

    def test_duplicate_clean_and_statistics(self):
        twice = pd.concat([self.financial, self.financial], ignore_index=True)
        cleaned, audit = clean_and_quality(twice, self.prices, self.news)
        self.assertEqual(audit.loc[0, '중복행수'], 1)
        self.assertEqual(len(cleaned['재무']), 1)
        stats = summarize_data(self.financial, self.prices, self.news)
        self.assertTrue({'재무', '주가', '뉴스'}.issubset(set(stats['구분'])))

    def test_fact_checks_and_export(self):
        df = verify_facts('테스트기업', self.financial, self.prices, self.news, '000000')
        self.assertEqual(df.loc[df['검증항목']=='2024년 자산=부채+자본', '상태'].iloc[0], '확인')
        self.assertGreater(len(excel_bytes(self.result)), 1000)
        context = make_ai_context(self.result)
        self.assertEqual(context['재무(억원)'][0]['매출액'], 0.0)

    def test_report_markdown(self):
        report = {k: '데이터에 근거한 정성 검토입니다.' for k in
                  ['profitability', 'stability', 'trend', 'news_risks', 'conclusion']}
        content = report_markdown(self.result, report)
        self.assertIn('테스트기업 투자검토 보고서', content)
        self.assertIn('사실 검증', content)

    def test_ai_numbers_blocked(self):
        fake_module = ModuleType('openai')
        client_mock = MagicMock()
        fake_module.OpenAI = client_mock
        fake = {'profitability':'수익 100', 'stability':'안정성 확인 필요',
                'trend':'추세 확인 필요', 'news_risks':'기사 진위 확인 필요', 'conclusion':'종합 확인 필요'}
        client_mock.return_value.chat.completions.create.return_value.choices = [
            type('choice', (), {'message':type('message', (), {'content':json.dumps(fake)})()})()]
        with patch.dict(sys.modules, {'openai': fake_module}):
            with self.assertRaisesRegex(ValueError, '숫자'):
                ai_review(self.result, 'test-not-real-key')



    @patch('core.fetch_public_json')
    def test_financial_api_stock_endpoint_and_fields(self, fake):
        from datetime import datetime
        from core import fetch_stock_public, STOCK_BASE
        fake.return_value = {'items': {'item': [
            {'basDt': '20261016', 'srtnCd': '005930', 'itmsNm': '삼성전자',
             'clpr': '70000', 'mkp': '69900', 'hipr': '70500', 'lopr': '69000',
             'trqu': '12345678', 'trPrc': '864000000000',
             'mrktTotAmt': '400000000000000', 'fltRt': '1.25'},
            # 005935 우선주가 검색에 잡혀도 정확한 코드가 아니므로 제외
            {'basDt': '20261016', 'srtnCd': '005935', 'itmsNm': '삼성전자우', 'clpr': '60000'},
        ]}, 'totalCount': 2}
        df, ticker = fetch_stock_public('삼성전자', 'fake-key', days=30,
                                        ticker='005930', today=datetime(2026, 10, 20))
        self.assertEqual(ticker, '005930')
        self.assertEqual(len(df), 1)
        self.assertEqual(df.loc[0, '종가'], 70000)
        self.assertEqual(df.loc[0, '거래량'], 12345678)
        self.assertEqual(df.loc[0, '시가총액'], 400000000000000)
        self.assertIn('금융위원회', df.loc[0, '출처'])
        args, kwargs = fake.call_args
        self.assertEqual(args[0], STOCK_BASE)
        self.assertEqual(args[1]['likeSrtnCd'], '005930')
        self.assertEqual(args[1]['endBasDt'], '20261021')
        self.assertEqual(args[1]['resultType'], 'json')

    @patch('core.fetch_public_json')
    def test_financial_api_stock_empty(self, fake):
        from core import fetch_stock_public
        fake.return_value = {'items': {}, 'totalCount': 0}
        df, code = fetch_stock_public('삼성전자', 'fake-key', ticker='005930')
        self.assertTrue(df.empty)
        self.assertEqual(code, '005930')

    @patch('core.fetch_stock_public')
    @patch('core.resolve_ticker')
    def test_public_stock_is_first_choice(self, mock_ticker, mock_public):
        from core import fetch_stock
        mock_ticker.return_value = ('005930', 'KS')
        official = pd.DataFrame([{'날짜': pd.Timestamp('2026-10-16'), '종가': 1000,
                                  '거래량': 50, '출처': '금융위원회 공식 시세'}])
        mock_public.return_value = (official, '005930')
        got, ticker, warns = fetch_stock('삼성전자', api_key='fake-key')
        self.assertEqual(ticker, '005930')
        self.assertEqual(len(got), 1)
        self.assertEqual(warns, [])
        mock_public.assert_called_once()

    @patch('core.requests.get')
    def test_public_api_http_error_never_leaks_key(self, fake_get):
        from core import fetch_public_json
        fake_get.return_value.status_code = 401
        try:
            fetch_public_json('https://example.com/api', {'serviceKey': 'secret-PROHIBITED'})
        except RuntimeError as exc:
            self.assertNotIn('secret-PROHIBITED', str(exc))
            self.assertIn('401', str(exc))
        else:
            self.fail('HTTP401 should raise')

    def test_fsc_columns_and_excel_sheets(self):
        from io import BytesIO
        from openpyxl import load_workbook
        data = excel_bytes(self.result)
        wb = load_workbook(BytesIO(data), read_only=True)
        try:
            self.assertTrue({'재무정보', '주가정보', '뉴스정보', '데이터품질',
                             '통계집계', '사실검증', '데이터출처', '수집경고', '차트'}.issubset(set(wb.sheetnames)))
        finally:
            wb.close()

    @patch('core.requests.get')
    def test_newsapi_priority(self, fake_get):
        from core import fetch_news
        fake_get.return_value.status_code = 200
        fake_get.return_value.json.return_value = {
            'status': 'ok',
            'articles': [{'title': '테스트 기사', 'url': 'https://example.com/test',
                          'description': '검색 미리보기', 'publishedAt': '2026-10-20',
                          'source': {'name': '예시신문'}}]}
        news = fetch_news('테스트기업', news_key='news-example', serpapi_key='serp-example')
        self.assertEqual(len(news), 1)
        self.assertEqual(news.iloc[0]['언론사'], '예시신문')
        self.assertEqual(news.iloc[0]['출처'], 'NewsAPI (기사 전문 아님)')
        self.assertIn('newsapi.org', fake_get.call_args.args[0])

    @patch('core.requests.get')
    def test_serpapi_fallback_and_no_credentials_in_export(self, fake_get):
        from core import fetch_news
        def fake_requests(url, **kwargs):
            response = MagicMock()
            if 'newsapi.org' in url:
                response.status_code = 401
            elif 'serpapi.com' in url:
                response.status_code = 200
                response.json.return_value = {
                    'news_results': [{'title': '테스트 SERP 기사', 'link': 'https://example.com/serp',
                                      'source': {'name': '예시'}, 'date': '2 hours ago',
                                      'snippet': '뉴스 검색 미리보기'}]}
            else:
                self.fail('SerpAPI success should not fall through to RSS')
            return response
        fake_get.side_effect = fake_requests
        news = fetch_news('테스트기업', news_key='NEWS-EXAMPLE', serpapi_key='SERP-EXAMPLE')
        self.assertEqual(news.iloc[0]['언론사'], '예시')
        self.assertIn('SerpAPI', news.iloc[0]['출처'])
        self.assertNotIn('NEWS-EXAMPLE', str(news.to_dict()))
        self.assertNotIn('SERP-EXAMPLE', str(news.to_dict()))

    def test_excel_contains_charts_when_data_exists(self):
        from io import BytesIO
        from openpyxl import load_workbook
        data = excel_bytes(self.result)
        wb = load_workbook(BytesIO(data))
        try:
            self.assertEqual(len(wb['차트']._charts), 3)
        finally:
            wb.close()

    @patch('core.fetch_public_json')
    def test_finance_percent_and_yoy_growth(self, fake):
        from datetime import datetime
        from core import fetch_financial
        def fake_json(url, params):
            year = int(params['bizYear'])
            return {'items': {'item': [{'crno': '1234567890123', 'bizYear': str(year),
                      'fnclDcdNm': '연결', 'basDt': str(year)+'1231',
                      'enpSaleAmt': str((year-2023)*1000), 'enpBzopPft': '200',
                      'enpCrtmNpf': '100', 'enpTastAmt':'10000',
                      'enpTdbtAmt':'4000', 'enpTcptAmt':'6000'}]}, 'totalCount': 1}
        fake.side_effect = fake_json
        df, warnings = fetch_financial('테스트기업', '1234567890123', 'FAKE', 2, datetime(2026, 10, 20))
        self.assertEqual(len(df), 2)
        self.assertEqual(df.loc[df['연도']==2025,'매출증감률(%)'].iloc[0], 100.0)
        self.assertEqual(df.loc[df['연도']==2025,'순이익률(%)'].iloc[0], 5.0)
        self.assertEqual(warnings, [])

    @patch('core.fetch_news')
    @patch('core.fetch_stock')
    @patch('core.fetch_financial')
    @patch('core.resolve_corporation')
    def test_end_to_end_flow_mocked(self, mock_corp, mock_fin, mock_stock, mock_news):
        from core import collect_all
        mock_corp.return_value = {'기업명': '테스트기업', '법인등록번호': '1234567890123',
                                  '출처': '테스트 기업 API'}
        mock_fin.return_value = (self.financial, [])
        mock_stock.return_value = (self.prices.assign(종목명='테스트기업'), '000000', [])
        mock_news.return_value = self.news
        result = collect_all('테스트기업', finance_key='FIN-MOCK', stock_key='STOCK-MOCK',
                             news_key='NEWS-MOCK', serpapi_key='SERP-MOCK')
        self.assertEqual(result['법인등록번호'], '1234567890123')
        self.assertEqual(result['종목코드'], '000000')
        self.assertEqual(len(result['재무']), 1)
        self.assertEqual(len(result['주가']), 1)
        self.assertEqual(len(result['뉴스']), 1)
        self.assertTrue({'차트', '재무정보'}.issubset(
            __import__('openpyxl').load_workbook(
                __import__('io').BytesIO(excel_bytes(result))).sheetnames))
        mock_news.assert_called_once_with('테스트기업', news_key='NEWS-MOCK',
                                          serpapi_key='SERP-MOCK')

if __name__ == '__main__':
    unittest.main()
